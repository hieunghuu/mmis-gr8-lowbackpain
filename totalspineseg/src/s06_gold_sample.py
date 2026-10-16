# -*- coding: utf-8 -*-
"""S06 - Bước 5a: tách 80 ghi chú chuẩn vàng (gold) + tạo file Excel để nhóm gán nhãn.

* Mỗi ghi chú xếp vào 1 nhóm bằng regex đơn giản, ĐỘC LẬP với bộ rule (ưu tiên từ trên
  xuống): uncommon (từ khóa bệnh ít gặp) > multi (>= 2 tầng) > single (1 tầng) > no_level.
* Lấy mẫu phân tầng THEO TỈ LỆ (seed 2026): 80 ghi chú chia đúng tỉ lệ từng nhóm.
* KHÔNG in nội dung ghi chú ra màn hình - giữ tập gold "chưa ai nhìn" khi viết rule (5b).

Đầu ra
  outputs/tables/5_nlp/nlp_split.csv        patient_id, stratum, split (dev / gold)
  data/gold/gold_annotation.xlsx      HuongDan / NhanBenh / QuyUoc / ViDu / Notes / Labels
                                      (+ sheet ẩn Lists cho danh sách thả xuống)

Chi tiết thiết kế: README mục 5.5 (quy ước) và 5.7 (chia dữ liệu).
Chạy (từ thư mục totalspineseg/):  python src/s06_gold_sample.py
"""
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import CODE, REPORT_XLSX, SEED, T_NLP

GOLD_DIR = CODE / "data" / "gold"
N_GOLD = 80

LEVEL_RE = re.compile(r"\b(?:t12|l[1-5])\s*[-/ ]\s*(?:l?[1-5]|s1?)\b", re.I)
UNCOMMON_RE = re.compile(
    r"listhesis|schmorl|modic|end.?plate|ha?emangioma|marrow|tarlov|fracture|wedg|sequest|"
    r"extru|facet|flavum|scolio|annul|tear|desicca|dehydrat|osteophyt|cyst", re.I)

LEVELS = ["T12-L1", "L1-L2", "L2-L3", "L3-L4", "L4-L5", "L5-S1", "UNSPEC"]
VERTEBRAE = ["", "T12", "L1", "L2", "L3", "L4", "L5", "S1"]
FINDINGS = ["disc_bulge", "disc_protrusion", "disc_herniation", "thecal_sac_compression",
            "nerve_root_compression", "foraminal_stenosis", "spinal_canal_stenosis",
            "annular_tear", "disc_desiccation", "muscle_spasm", "spondylolisthesis",
            "endplate_degeneration", "schmorl_nodes", "bone_marrow_lesion",
            "ligamentum_flavum_hypertrophy", "facet_arthropathy", "scoliosis",
            "tarlov_cyst", "vertebral_fracture", "osteophyte", "disc_space_narrowing",
            "other", "none"]
PRESENT = ["TRUE", "FALSE"]
SEVERITY = ["unspecified", "mild", "moderate", "severe"]
SIDE = ["unspecified", "left", "right", "bilateral", "central"]
COLS = ["patient_id", "level", "vertebra", "root", "finding", "present", "severity", "side", "comment"]

# Ví dụ gán nhãn mẫu - lấy ghi chú đầu tiên trong danh sách thuộc TẬP DEV
EXAMPLES = {
    4: [("UNSPEC", "", "", "muscle_spasm", "TRUE", "unspecified", "unspecified", "Q7: câu không nhắc tầng"),
        ("L4-L5", "", "", "disc_bulge", "TRUE", "unspecified", "unspecified", "Một câu nhắc 2 tầng → mỗi tầng 1 dòng"),
        ("L4-L5", "", "", "thecal_sac_compression", "TRUE", "mild", "unspecified", "\"mildly compressing\" → mild"),
        ("L4-L5", "", "", "nerve_root_compression", "TRUE", "unspecified", "bilateral", "Q2 \"encroaching\"; Q11 \"roots\" số nhiều"),
        ("L5-S1", "", "", "disc_bulge", "TRUE", "unspecified", "unspecified", ""),
        ("L5-S1", "", "", "thecal_sac_compression", "TRUE", "mild", "unspecified", ""),
        ("L5-S1", "", "", "nerve_root_compression", "TRUE", "unspecified", "bilateral", "")],
    1: [("L4-L5", "", "", "disc_bulge", "TRUE", "unspecified", "left", "\"more to the left side\""),
        ("L4-L5", "", "", "thecal_sac_compression", "TRUE", "unspecified", "unspecified", ""),
        ("L4-L5", "", "", "nerve_root_compression", "TRUE", "unspecified", "left", "\"left nerve root\""),
        ("L4-L5", "", "", "foraminal_stenosis", "TRUE", "unspecified", "right", "\"narrowing right neural foramen\""),
        ("L4-L5", "", "", "annular_tear", "TRUE", "unspecified", "left", "Kế thừa tầng L4-L5 từ câu trước")],
}

# (nhãn, tên tiếng Việt, cách bác sĩ thường viết, ghi chú)
LABEL_TABLE = [
    ("disc_bulge", "Phồng đĩa đệm", "disc bulge, diffuse bulge, bulging", ""),
    ("disc_protrusion", "Lồi đĩa đệm", "protrusion, protruding", "Lồi khu trú; có thể đi kèm phồng (Q4)"),
    ("disc_herniation", "Thoát vị đĩa đệm", "herniation; extruded / sequestrated / detached disc fragment",
     "Chữ \"herniation\" thường nằm trong câu phủ định. Mảnh đĩa tách rời / đẩy ra = thoát vị (Q3)"),
    ("thecal_sac_compression", "Chèn ép bao màng cứng", "compressing / indenting the thecal sac",
     "\"abutting\" (chạm) KHÔNG phải chèn ép (Q1)"),
    ("nerve_root_compression", "Chèn ép rễ thần kinh", "compressing / encroaching (exit) nerve root(s)",
     "\"encroaching\" = chèn ép (Q2). \"Lt S1 nerve root\": ghi S1 vào cột root"),
    ("foraminal_stenosis", "Hẹp lỗ liên hợp", "narrowing neural foramen; encroaching exit (neural) canal",
     "Exit canal = lỗ liên hợp (Q5)"),
    ("spinal_canal_stenosis", "Hẹp ống sống", "canal stenosis, narrowed (spinal) canal",
     "\"adequate spinal canal\" → FALSE (Q6). \"extending to lateral recess\" chỉ là vị trí, không phải hẹp"),
    ("annular_tear", "Rách vòng xơ", "annular tear, torn annulus, high signal in the annulus", ""),
    ("disc_desiccation", "Mất nước đĩa đệm", "desiccation, dehydrated disc; \"dissicating disc material\" (viết sai)", ""),
    ("muscle_spasm", "Co cứng cơ", "muscle spasm", "Luôn ghi tầng UNSPEC"),
    ("spondylolisthesis", "Trượt đốt sống", "spondylolisthesis, listhesis, \"L5 over S1\"",
     "\"L5 over S1\" → tầng L5-S1; độ I/II ghi vào comment"),
    ("endplate_degeneration", "Thoái hóa tấm tận", "Modic changes, endplate changes / irregular endplates", ""),
    ("schmorl_nodes", "Nốt Schmorl", "Schmorl's node(s)", "Thường gắn với đốt → cột vertebra"),
    ("bone_marrow_lesion", "Tổn thương tủy xương", "marrow edema / signal change, hemangioma", "Thường gắn với đốt → cột vertebra"),
    ("ligamentum_flavum_hypertrophy", "Phì đại dây chằng vàng", "ligamentum flavum hypertrophy", ""),
    ("facet_arthropathy", "Thoái hóa khớp mặt", "facet arthropathy / hypertrophy / degeneration", ""),
    ("scoliosis", "Vẹo cột sống", "scoliosis", "Luôn ghi tầng UNSPEC"),
    ("tarlov_cyst", "Nang Tarlov", "Tarlov cyst, perineural cyst", ""),
    ("vertebral_fracture", "Gãy / lún đốt sống", "fracture, wedging, collapse", "Thường gắn với đốt → cột vertebra"),
    ("osteophyte", "Gai xương", "osteophyte(s), spur", ""),
    ("disc_space_narrowing", "Hẹp khoang đĩa", "narrowed disc space(s)", "Đĩa bị xẹp – đo được trên ảnh"),
    ("other", "Phát hiện cột sống khác", "vd hậu phẫu, \"degenerative spondylotic changes\"", "Bắt buộc mô tả ở cột comment"),
    ("none", "Không có phát hiện nào", "–", "Ghi chú không nhắc phát hiện nào → 1 dòng none"),
]

# (mã, gặp câu như thế nào, gán thế nào, lý do)
RULES = [
    ("Q1", "\"abutting\" (chạm) thecal sac / nerve root", "Ghi chèn ép với present = FALSE", "Chạm ≠ đè"),
    ("Q2", "\"encroaching (exit) nerve roots\"", "nerve_root_compression = TRUE", "Bác sĩ dùng thay cho \"compressing\""),
    ("Q3", "\"sequestrated / extruded / detached disc fragment\"", "disc_herniation = TRUE", "Theo danh pháp Fardon đây là dạng thoát vị"),
    ("Q4", "\"protrusion on top of disc bulge\"", "Ghi CẢ HAI: disc_protrusion và disc_bulge", "Hai hình thái cùng tồn tại"),
    ("Q5", "\"compressing / encroaching exit (neural) canal\"", "foraminal_stenosis = TRUE", "Exit canal = lỗ liên hợp"),
    ("Q6", "\"adequate spinal canal\"", "spinal_canal_stenosis, present = FALSE", "Câu mô tả bình thường"),
    ("Q7", "Câu không nhắc tầng nào", "level = UNSPEC", "Không ghép được với ảnh theo tầng"),
    ("Q8", "\"multilevel … noticeably at L3-L4 & L4-L5\"", "Chỉ ghi các tầng được nêu tên", "Không suy diễn tầng không được viết"),
    ("Q9", "Tầng chỉ dùng để định vị (\"below L4/L5 level\")", "Ghi tầng đĩa được nêu", ""),
    ("Q10", "Tầng không được nhắc", "Không ghi gì", "Không tự suy ra là bình thường"),
    ("Q11", "\"nerve roots\" / \"exit nerve roots\" / \"exit (neural) canals\" (số nhiều), không nói bên",
     "side = bilateral", "Số nhiều = hai bên; số ít (\"exit canal\") → unspecified"),
    ("Q12", "Câu phủ định khuôn mẫu KHÔNG có tầng (vd \"No significant thecal sac or nerve root compression\", "
            "\"adequate spinal canal\") đứng sau câu có tầng",
     "level = UNSPEC, không lấy tầng của câu trước", "Câu khuôn mẫu nói chung cho cả phim"),
]
SEVERITY_TABLE = [("mild", "mild, mildly, minimal, small, slight"), ("moderate", "moderate"),
                  ("severe", "severe, significant(ly), marked, largely"),
                  ("unspecified", "Không có từ chỉ mức độ, hoặc dòng FALSE")]
SIDE_TABLE = [("left", "left, Lt"), ("right", "right, Rt"), ("bilateral", "both, bilateral; \"nerve roots\", \"exit canals\" số nhiều (Q11)"),
              ("central", "central, median"), ("unspecified", "Không nói bên, hoặc dòng FALSE")]

COL_HELP = [  # (cột, ý nghĩa, giá trị cho phép, ví dụ)
    ("patient_id", "Mã bệnh nhân (chọn trong danh sách 80 mã ở sheet Notes)", "80 mã gold", "15"),
    ("level", "Tầng đĩa đệm của phát hiện", "T12-L1 … L5-S1, UNSPEC", "L4-L5"),
    ("vertebra", "Chỉ khi phát hiện gắn với một ĐỐT (vd \"L2 vertebral body\")", "T12, L1 … S1 hoặc để trống", "L2"),
    ("root", "Tên rễ thần kinh nếu được nêu (vd \"Lt S1 nerve root\")", "Gõ tay, vd S1", "S1"),
    ("finding", "Loại phát hiện – tra bảng ở sheet NhanBenh", "23 nhãn", "disc_bulge"),
    ("present", "TRUE = có; FALSE = bị phủ định (\"No evidence of …\")", "TRUE / FALSE", "TRUE"),
    ("severity", "Mức độ – tra bảng ở sheet QuyUoc", "unspecified, mild, moderate, severe", "mild"),
    ("side", "Bên – tra bảng ở sheet QuyUoc", "unspecified, left, right, bilateral, central", "left"),
    ("comment", "Ghi chú khi phân vân, hoặc mô tả cho nhãn other", "Tự do, ngắn", "không chắc tầng"),
]

HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
SECTION_FONT = Font(bold=True, size=12, color="1F4E79")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def _title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=15, color="1F4E79")
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(italic=True, color="595959")
    return 4


def _section(ws, row, text):
    ws.cell(row=row, column=1, value=text).font = SECTION_FONT
    return row + 1


def _table(ws, row, header, rows, height=None):
    for j, h in enumerate(header, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font, c.fill, c.alignment, c.border = Font(bold=True, color="FFFFFF"), HEAD_FILL, WRAP, BOX
    for r in rows:
        row += 1
        for j, v in enumerate(r, 1):
            c = ws.cell(row=row, column=j, value=v)
            c.alignment, c.border = WRAP, BOX
        if height:
            ws.row_dimensions[row].height = height
    return row + 2


def _widths(ws, widths):
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def sheet_guide(wb, n_gold):
    ws = wb.active
    ws.title = "HuongDan"
    r = _title(ws, f"Hướng dẫn gán nhãn bộ chuẩn vàng – {n_gold} ghi chú bác sĩ",
               "Bộ nhãn này là ĐÁP ÁN để chấm độ chính xác của bộ rule NLP (README bước 5).")
    r = _section(ws, r, "1. Các sheet trong file")
    r = _table(ws, r, ["Sheet", "Dùng để"], [
        ("HuongDan", "Đọc trước: các bước và ý nghĩa từng cột"),
        ("NhanBenh", "Bảng tra 23 nhãn: tên tiếng Việt, cách bác sĩ thường viết"),
        ("QuyUoc", "Các tình huống khó (Q1–Q12), cách chọn mức độ và bên"),
        ("ViDu", "Một ghi chú đã gán nhãn mẫu (không thuộc 80 ghi chú cần gán)"),
        ("Notes", f"{n_gold} ghi chú cần gán; cột \"da_gan\" đếm số dòng đã nhập (đỏ = chưa gán)"),
        ("Labels", "NƠI NHẬP NHÃN – mỗi phát hiện một dòng, chọn giá trị từ danh sách thả xuống"),
    ])
    r = _section(ws, r, "2. Các bước")
    steps = [
        "Đọc sheet NhanBenh, QuyUoc và ví dụ ở sheet ViDu (khoảng 10 phút).",
        "Ở sheet Notes, đọc lần lượt từng ghi chú.",
        "Mỗi PHÁT HIỆN được nhắc → nhập 1 dòng ở sheet Labels. Một câu nhắc 2 tầng → 2 dòng.",
        "Ghi CẢ phát hiện bị phủ định (vd \"No evidence of disc herniation\") với present = FALSE.",
        "Ghi chú không có phát hiện nào → 1 dòng finding = none, level = UNSPEC.",
        "Phân vân → vẫn ghi theo cách hiểu tốt nhất và giải thích ngắn ở cột comment.",
        "Xong: kiểm tra sheet Notes – không còn ô đỏ ở cột da_gan.",
    ]
    r = _table(ws, r, ["Bước", "Việc cần làm"], [(i, s) for i, s in enumerate(steps, 1)])
    r = _section(ws, r, "3. Ý nghĩa các cột ở sheet Labels")
    r = _table(ws, r, ["Cột", "Ý nghĩa", "Giá trị cho phép", "Ví dụ"], COL_HELP)
    r = _section(ws, r, "4. Gán tầng (level)")
    r = _table(ws, r, ["Tình huống", "Gán thế nào"], [
        ("Câu có tầng (L4/5, L4-5, L4-L5 …)", "Ghi đúng tầng đó, viết theo dạng chuẩn L4-L5, L5-S1"),
        ("Câu KHẲNG ĐỊNH không có tầng nhưng nói tiếp câu trước có tầng", "Dùng tầng của câu trước"),
        ("Câu PHỦ ĐỊNH khuôn mẫu không có tầng (\"No significant … compression\")", "UNSPEC (Q12)"),
        ("Không có tầng nào", "UNSPEC (Q7)"),
        ("\"Lt S1 nerve root\"", "S1 là tên rễ → cột root, KHÔNG phải tầng"),
        ("\"L2 vertebral body\"", "Cột vertebra = L2, level = UNSPEC (trừ khi câu có tầng đĩa)"),
    ])
    r = _section(ws, r, "5. Làm độc lập")
    ws.cell(row=r, column=1, value="Hai người gán riêng (copy thành gold_labels_A.xlsx, gold_labels_B.xlsx), không trao đổi "
            "trước khi xong; sau đó so sánh, thống nhất thành gold_labels_final.xlsx.").alignment = WRAP
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    ws.row_dimensions[r].height = 32
    _widths(ws, [30, 62, 34, 16])


def sheet_labels_ref(wb):
    ws = wb.create_sheet("NhanBenh")
    r = _title(ws, "Bảng tra nhãn (cột finding)", "Tra cột \"Cách bác sĩ thường viết\" khi gặp cụm từ trong ghi chú.")
    _table(ws, r, ["Nhãn (finding)", "Tên tiếng Việt", "Cách bác sĩ thường viết", "Ghi chú"], LABEL_TABLE)
    _widths(ws, [30, 24, 48, 58])
    ws.freeze_panes = f"A{r + 1}"


def sheet_rules(wb):
    ws = wb.create_sheet("QuyUoc")
    r = _title(ws, "Quy ước cho các tình huống khó", "Giống README mục 5.5 – dùng chung cho người gán nhãn và bộ rule.")
    r = _table(ws, r, ["Mã", "Gặp câu như thế nào", "Gán thế nào", "Lý do"], RULES)
    r = _section(ws, r, "Mức độ (severity) – chỉ cho dòng present = TRUE")
    r = _table(ws, r, ["Giá trị", "Từ trong ghi chú"], SEVERITY_TABLE)
    r = _section(ws, r, "Bên (side) – chỉ cho dòng present = TRUE")
    r = _table(ws, r, ["Giá trị", "Từ trong ghi chú"], SIDE_TABLE)
    ws.cell(row=r, column=1, value="Chỉ dùng từ gắn với CHÍNH phát hiện đó: \"disc protrusion, largely compressing the "
            "thecal sac\" → \"largely\" thuộc về chèn ép, không thuộc về protrusion.").alignment = WRAP
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    ws.row_dimensions[r].height = 32
    _widths(ws, [12, 46, 40, 40])


def sheet_example(wb, pid, note):
    ws = wb.create_sheet("ViDu")
    r = _title(ws, f"Ví dụ gán nhãn – bệnh nhân {pid}",
               "Ghi chú này thuộc tập phát triển, KHÔNG nằm trong danh sách cần gán.")
    ws.cell(row=r, column=1, value="Ghi chú gốc:").font = Font(bold=True)
    ws.cell(row=r + 1, column=1, value=note).alignment = WRAP
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=9)
    ws.row_dimensions[r + 1].height = 48
    _table(ws, r + 3, COLS[:-1] + ["giải thích"], [(pid, *x) for x in EXAMPLES[pid]])
    _widths(ws, [11, 9, 9, 7, 26, 9, 12, 12, 40])


def sheet_notes(wb, gold):
    ws = wb.create_sheet("Notes")
    _table(ws, 1, ["patient_id", "note", "da_gan"], [])
    for i, (pid, note) in enumerate(gold[["patient_id", "note"]].itertuples(index=False), start=2):
        ws.cell(row=i, column=1, value=int(pid)).alignment = WRAP
        ws.cell(row=i, column=2, value=note).alignment = WRAP
        ws.cell(row=i, column=3, value=f"=COUNTIF(Labels!$A:$A,A{i})").alignment = Alignment(horizontal="center", vertical="top")
        for j in (1, 2, 3):
            ws.cell(row=i, column=j).border = BOX
        n_lines = sum(max(1, -(-len(part) // 105)) for part in str(note).splitlines() or [""])
        ws.row_dimensions[i].height = max(18, 15 * n_lines)
    last = len(gold) + 1
    ws.conditional_formatting.add(f"C2:C{last}", CellIsRule(operator="equal", formula=["0"],
                                  fill=PatternFill("solid", fgColor="F8CBAD")))
    ws.conditional_formatting.add(f"C2:C{last}", CellIsRule(operator="greaterThan", formula=["0"],
                                  fill=PatternFill("solid", fgColor="C6EFCE")))
    _widths(ws, [11, 110, 9])
    ws.freeze_panes = "A2"


def sheet_lists(wb):
    lists = wb.create_sheet("Lists")
    for j, vals in enumerate([LEVELS, VERTEBRAE, FINDINGS, PRESENT, SEVERITY, SIDE]):
        for i, v in enumerate(vals, 1):
            lists.cell(row=i, column=j + 1, value=v)
    lists.sheet_state = "hidden"


def sheet_input(wb, n_gold):
    ws = wb.create_sheet("Labels")
    _table(ws, 1, COLS, [])
    _widths(ws, [11, 10, 10, 7, 30, 9, 12, 12, 45])
    ws.freeze_panes = "A2"
    n_rows = 1000
    help_by_col = {c: h for c, h, _, _ in COL_HELP}
    specs = [("A", None), ("B", ("A", len(LEVELS))), ("C", ("B", len(VERTEBRAE))),
             ("E", ("C", len(FINDINGS))), ("F", ("D", len(PRESENT))),
             ("G", ("E", len(SEVERITY))), ("H", ("F", len(SIDE)))]
    for col, src in specs:
        if src is None:   # patient_id: chỉ cho chọn đúng các mã ở sheet Notes
            formula = f"=Notes!$A$2:$A${n_gold + 1}"
        else:
            lc, n = src
            formula = f"=Lists!${lc}$1:${lc}${n}"
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        name = COLS["ABCDEFGHI".index(col)]
        dv.promptTitle, dv.prompt = name, help_by_col[name][:250]
        dv.errorTitle, dv.error = "Giá trị không hợp lệ", "Hãy chọn từ danh sách thả xuống."
        dv.showInputMessage = True
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{n_rows}")


def stratify(notes):
    n_lv = notes.apply(lambda s: len({m.lower().replace(" ", "") for m in LEVEL_RE.findall(s)}))
    st = pd.Series("single", index=notes.index)
    st[n_lv == 0] = "no_level"
    st[n_lv >= 2] = "multi"
    st[notes.str.contains(UNCOMMON_RE)] = "uncommon"
    return st


def allocate(sizes, n):
    """Chia n theo tỉ lệ (phương pháp số dư lớn nhất)."""
    raw = sizes / sizes.sum() * n
    alloc = np.floor(raw).astype(int)
    for k in (raw - alloc).sort_values(ascending=False).index[: n - alloc.sum()]:
        alloc[k] += 1
    return alloc


def build_workbook(gold, example_pid, example_note):
    wb = Workbook()
    sheet_guide(wb, len(gold))
    sheet_labels_ref(wb)
    sheet_rules(wb)
    sheet_example(wb, example_pid, example_note)
    sheet_notes(wb, gold)
    sheet_input(wb, len(gold))
    sheet_lists(wb)
    return wb


def main():
    d = pd.read_excel(REPORT_XLSX)
    d.columns = ["patient_id", "note"]
    d = d.dropna(subset=["note"]).reset_index(drop=True)
    d["patient_id"] = d.patient_id.astype(int)
    d["stratum"] = stratify(d.note.astype(str))

    sizes = d.stratum.value_counts()
    alloc = allocate(sizes, N_GOLD)
    rng = np.random.default_rng(SEED)
    gold_ids = []
    for s, k in alloc.items():
        ids = d.loc[d.stratum == s, "patient_id"].to_numpy()
        gold_ids += rng.choice(ids, size=k, replace=False).tolist()
    d["split"] = np.where(d.patient_id.isin(gold_ids), "gold", "dev")
    d[["patient_id", "stratum", "split"]].sort_values("patient_id").to_csv(
        T_NLP / "nlp_split.csv", index=False)

    example_pid = next(p for p in EXAMPLES if p in set(d.loc[d.split == "dev", "patient_id"]))
    example_note = d.loc[d.patient_id == example_pid, "note"].iloc[0]
    gold = d[d.split == "gold"].sort_values("patient_id")
    GOLD_DIR.mkdir(parents=True, exist_ok=True)
    out = GOLD_DIR / "gold_annotation.xlsx"
    build_workbook(gold, example_pid, example_note).save(out)

    print("=== Bước 5a ===")
    print(pd.DataFrame({"số_ghi_chú": sizes, "lấy_vào_gold": alloc}).to_string())
    print(f"dev: {(d.split == 'dev').sum()} | gold: {(d.split == 'gold').sum()}")
    neg = d.note.str.contains(r"\bno\b|without|\bnot\b", case=False, regex=True)
    print(f"Gold có câu phủ định: {int((neg & (d.split == 'gold')).sum())} / {N_GOLD}")
    print(f"Ví dụ mẫu: BN {example_pid} (tập dev)")
    print(f"Đã ghi {T_NLP / 'nlp_split.csv'}")
    print(f"Đã ghi {out}")


if __name__ == "__main__":
    main()
